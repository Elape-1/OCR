from __future__ import annotations

import logging
import os
import shutil
import subprocess
import tempfile
import time
import threading
from datetime import datetime, timezone
from pathlib import Path
import uuid
from typing import Any, List
from io import BytesIO

import fitz
from fastapi import APIRouter, Depends, File, HTTPException, Query, UploadFile
from fastapi.responses import Response
from PIL import Image, ImageDraw, ImageFont, ImageSequence
from docx import Document as DocxDocument
from docx.document import Document as DocxDocumentPart
from docx.table import Table
from docx.text.paragraph import Paragraph
from docx.oxml.ns import qn
from sqlalchemy import String, cast, or_, select
from sqlalchemy.orm import Session, selectinload

from app.db import SessionLocal, get_db, set_rls_user
from app.auth import CurrentUserId, owner_scope
from app.document_schemas import classify_document_type
from app.model_inference import LayoutLMv3InferenceService
from app.models import Attribute, Document as DocumentRecord, Page
from app.ocr_processor import extract_tokens_with_bboxes
from app.routing_engine import build_attribute_payloads
from app.tasks import celery_app
from app.storage import delete_file, download_file, storage_enabled, upload_file

logger = logging.getLogger(__name__)
router = APIRouter()

SUPPORTED_FORMATS = {
    ".pdf": {"application/pdf"},
    ".docx": {
        "application/vnd.openxmlformats-officedocument.wordprocessingml.document",
    },
    ".doc": {"application/msword"},
    ".jpeg": {"image/jpeg"},
    ".jpg": {"image/jpeg"},
    ".png": {"image/png"},
    ".tiff": {"image/tiff"},
    ".tif": {"image/tiff"},
    ".bmp": {"image/bmp"},
}
IMAGE_EXTENSIONS = {".jpeg", ".jpg", ".png", ".tiff", ".tif", ".bmp"}
UPLOAD_ROOT = Path(os.getenv("UPLOAD_DIR", "uploads"))
RAW_DATASET_ROOT = Path(os.getenv("RAW_DATASET_PAGES_DIR", "datasets/raw-pages"))
POPPLER_PATH = os.getenv("POPPLER_PATH")
SOFFICE_CMD = os.getenv("SOFFICE_CMD", "soffice")
DOCUMENT_STATUSES = frozenset({"PENDING", "QUEUED", "PROCESSING", "PROCESSED", "FAILED", "APPROVED", "CANCELED"})
DOCUMENT_STATUS_TRANSITIONS = {
    "PENDING": {"QUEUED", "CANCELED"},
    "QUEUED": {"PROCESSING", "CANCELED"},
    "PROCESSING": {"PROCESSED", "FAILED", "CANCELED"},
    "PROCESSED": {"APPROVED", "PROCESSING"},
    "FAILED": {"QUEUED"},
    "APPROVED": {"PROCESSING"},
    "CANCELED": {"QUEUED"},
}


def _transition_document_status(document: DocumentRecord, target_status: str) -> None:
    target = target_status.strip().upper()
    current = str(document.status or "").upper()
    if target not in DOCUMENT_STATUSES:
        raise ValueError(f"Unsupported document status: {target_status}")
    if target != current and target not in DOCUMENT_STATUS_TRANSITIONS.get(current, set()):
        raise ValueError(f"Invalid document status transition: {current} -> {target}")
    if target == "PROCESSING" and document.processing_started_at is None:
        raise ValueError("PROCESSING requires processing_started_at")
    if target == "PROCESSED":
        if not document.pages or document.page_count != len(document.pages):
            raise ValueError("PROCESSED requires all document pages")
        if document.processing_finished_at is None or document.processing_duration_seconds is None:
            raise ValueError("PROCESSED requires processing timing data")
    if target == "FAILED" and (
        document.processing_finished_at is None or document.processing_duration_seconds is None
    ):
        raise ValueError("FAILED requires processing timing data")
    if target == "APPROVED" and not document.pages:
        raise ValueError("APPROVED requires at least one document page")
    if target == current:
        return
    document.status = target


@router.post("/ingest/validate")
async def validate_upload(file: UploadFile = File(...), user_id: CurrentUserId = None) -> dict[str, Any]:
    extension = Path(file.filename or "").suffix.lower()
    content_type = (file.content_type or "").lower()
    if extension not in SUPPORTED_FORMATS:
        raise HTTPException(status_code=400, detail=f"Unsupported file extension: {extension or 'missing'}")

    expected_types = SUPPORTED_FORMATS[extension]
    if content_type not in expected_types:
        raise HTTPException(
            status_code=400,
            detail=f"MIME type {content_type or 'missing'} does not match supported types for {extension}",
        )

    return {
        "valid": True,
        "filename": file.filename,
        "extension": extension,
        "content_type": content_type,
    }


@router.post("/ingest")
async def ingest_document(file: UploadFile = File(...), db: Session = Depends(get_db), user_id: CurrentUserId = None) -> dict[str, Any]:
    validation = await validate_upload(file)
    UPLOAD_ROOT.mkdir(parents=True, exist_ok=True)

    extension = Path(file.filename or "document").suffix.lower()
    document = DocumentRecord(
        name=file.filename or "document",
        format=extension.lstrip("."),
        owner_id=user_id,
        page_count=0,
        status="QUEUED",
    )
    db.add(document)
    db.flush()
    document_id = document.id

    staging_dir = UPLOAD_ROOT / str(document_id)
    staging_dir.mkdir(parents=True, exist_ok=True)

    stored_path = staging_dir / f"source{extension}"
    source_key = f"{user_id}/{document_id}/source{extension}"
    try:
        with stored_path.open("wb") as target:
            shutil.copyfileobj(file.file, target)
        upload_file(stored_path, source_key)
        document.status = "QUEUED"
        db.commit()
        db.refresh(document)
    except Exception as exc:
        db.rollback()
        logger.exception("Failed to persist uploaded file %s", file.filename)
        raise HTTPException(status_code=500, detail="Failed to store uploaded document") from exc

    task_id = None
    if celery_app.conf.task_always_eager:
        thread = threading.Thread(
            target=process_document_job.run,
            args=(document_id, source_key, file.filename or "document", user_id),
            daemon=True,
        )
        thread.start()
    else:
        try:
            task = process_document_job.delay(document_id, source_key, file.filename or "document", user_id)
            task_id = task.id
        except Exception as exc:
            logger.warning(
                "Celery task dispatch failed; falling back to background thread for %s: %s",
                file.filename,
                exc,
            )
            thread = threading.Thread(
                target=process_document_job.run,
                args=(document_id, source_key, file.filename or "document", user_id),
                daemon=True,
            )
            thread.start()

    return {
        "job_id": task_id,
        "document_id": document_id,
        "validation": validation,
        "status": "QUEUED",
    }


@router.post("/documents/{document_id}/reprocess")
def reprocess_document(document_id: int, db: Session = Depends(get_db), user_id: CurrentUserId = None) -> dict[str, Any]:
    document = db.query(DocumentRecord).filter(
        DocumentRecord.id == document_id,
        owner_scope(DocumentRecord.owner_id, user_id),
    ).first()
    if document is None:
        raise HTTPException(status_code=404, detail="Document not found")
    if document.status not in {"CANCELED", "FAILED"}:
        raise HTTPException(status_code=400, detail="Only canceled or failed documents can be processed again")

    extension = f".{document.format.lower().lstrip('.') }"
    source_key = f"{document.owner_id}/{document.id}/source{extension}"
    _transition_document_status(document, "QUEUED")
    db.commit()

    task_id = None
    if celery_app.conf.task_always_eager:
        thread = threading.Thread(
            target=process_document_job.run,
            args=(document.id, source_key, document.name, user_id),
            daemon=True,
        )
        thread.start()
    else:
        task = process_document_job.delay(document.id, source_key, document.name, user_id)
        task_id = task.id

    return {"document_id": document.id, "job_id": task_id, "status": "QUEUED"}


@router.get("/documents/{document_id}")
def get_document(document_id: int, db: Session = Depends(get_db), user_id: CurrentUserId = None) -> dict[str, Any]:
    document = db.query(DocumentRecord).filter(DocumentRecord.id == document_id, owner_scope(DocumentRecord.owner_id, user_id)).first()
    if document is None:
        raise HTTPException(status_code=404, detail="Document not found")

    return {
        "id": document.id,
        "name": document.name,
        "format": document.format,
        "document_type": document.document_type,
        "page_count": document.page_count,
        "timestamp": document.timestamp,
        "status": document.status,
        "pages": [
            {
                "id": page.id,
                "page_number": page.page_number,
                "image_path": page.image_path,
                "image_url": f"/api/v1/documents/{document.id}/pages/{page.page_number}/image",
            }
            for page in document.pages
        ],
    }


@router.get("/documents/{document_id}/attributes")
def get_document_attributes(document_id: int, db: Session = Depends(get_db), user_id: CurrentUserId = None) -> dict[str, Any]:
    document = db.query(DocumentRecord).filter(DocumentRecord.id == document_id, owner_scope(DocumentRecord.owner_id, user_id)).first()
    if document is None:
        raise HTTPException(status_code=404, detail="Document not found")

    saved_attribute_rows = []
    for page in document.pages:
        saved_attribute_rows.extend(
            [
                {
                    "id": attribute.id,
                    "page_id": attribute.page_id,
                    "document_id": attribute.document_id,
                    "entity_type_label": attribute.entity_type_label,
                    "extracted_value": attribute.extracted_value,
                    "original_prediction": attribute.extracted_value,
                    "bounding_boxes": attribute.bounding_boxes,
                    "confidence_score": attribute.confidence_score,
                    "validation_status": attribute.validation_status,
                    "source": attribute.source,
                    "saved": True,
                }
                for attribute in page.attributes
            ]
        )

    staged_attribute_rows = [
        {
            "temp_id": item.get("temp_id") or f"staged-{index}",
            "page_id": item.get("page_id"),
            "page_number": item.get("page_number"),
            "document_id": document.id,
            "entity_type_label": item.get("entity_type_label", ""),
            "extracted_value": item.get("extracted_value", ""),
            "original_prediction": item.get("original_prediction", item.get("extracted_value", "")),
            "bounding_boxes": item.get("bounding_boxes", []),
            "confidence_score": item.get("confidence_score", 0.0),
            "validation_status": item.get("validation_status", "PENDING"),
            "source": item.get("source", "model"),
            "saved": False,
        }
        for index, item in enumerate(document.extracted_attributes or [], start=1)
        if isinstance(item, dict)
    ]

    # During draft review, prioritize staged attributes so unapproved items remain actionable.
    attribute_rows = staged_attribute_rows if staged_attribute_rows else saved_attribute_rows

    return {
        "document_id": document.id,
        "review_state": "draft" if staged_attribute_rows and not saved_attribute_rows else "saved",
        "attributes": attribute_rows,
    }


@router.get("/documents")
def list_documents(q: str | None = Query(None), limit: int = Query(50, ge=1, le=200), db: Session = Depends(get_db), user_id: CurrentUserId = None) -> dict[str, Any]:
    query = (q or "").strip().lower()
    statement = db.query(DocumentRecord).options(
        selectinload(DocumentRecord.pages).selectinload(Page.attributes)
    ).filter(owner_scope(DocumentRecord.owner_id, user_id))
    if query:
        like = f"%{query}%"
        attribute_match = select(Attribute.id).where(
            Attribute.document_id == DocumentRecord.id,
            or_(
                Attribute.entity_type_label.ilike(like),
                Attribute.extracted_value.ilike(like),
                Attribute.validation_status.ilike(like),
            ),
        ).exists()
        statement = statement.filter(
            or_(
                DocumentRecord.name.ilike(like),
                DocumentRecord.format.ilike(like),
                DocumentRecord.document_type.ilike(like),
                DocumentRecord.status.ilike(like),
                cast(DocumentRecord.page_count, String).ilike(like),
                cast(DocumentRecord.timestamp, String).ilike(like),
                attribute_match,
            )
        )
    documents = statement.order_by(DocumentRecord.timestamp.desc()).limit(limit).all()

    summaries: list[dict[str, Any]] = []
    for document in documents:
        attribute_summaries: list[dict[str, Any]] = []
        for page in sorted(document.pages, key=lambda item: item.page_number):
            for attribute in page.attributes:
                attribute_summaries.append(
                    {
                        "id": attribute.id,
                        "label": attribute.entity_type_label,
                        "value": attribute.extracted_value,
                        "validation_status": attribute.validation_status,
                        "confidence_score": attribute.confidence_score,
                    }
                )

        summaries.append(
            {
                "id": document.id,
                "name": document.name,
                "format": document.format,
                "document_type": document.document_type,
                "page_count": document.page_count,
                "timestamp": document.timestamp,
                "status": document.status,
                "attribute_count": len(attribute_summaries),
                "attribute_summaries": attribute_summaries[:6],
                "search_text": " ".join(
                    [
                        document.name,
                        document.format,
                        document.document_type,
                        document.status,
                        " ".join(summary["label"] for summary in attribute_summaries),
                        " ".join(summary["value"] for summary in attribute_summaries),
                    ]
                ),
            }
        )
    return {"documents": summaries, "query": q or ""}


@router.get("/documents/{document_id}/ocr")
def get_document_ocr(document_id: int, db: Session = Depends(get_db), user_id: CurrentUserId = None) -> dict[str, Any]:
    document = db.query(DocumentRecord).filter(DocumentRecord.id == document_id, owner_scope(DocumentRecord.owner_id, user_id)).first()
    if document is None:
        raise HTTPException(status_code=404, detail="Document not found")

    page_payloads: list[dict[str, Any]] = []
    for page in sorted(document.pages, key=lambda item: item.page_number):
        image_path = Path(page.image_path)
        if not image_path.exists():
            image_path = RAW_DATASET_ROOT / str(document.id) / f"page-{page.page_number:03d}.png"
        if not image_path.exists():
            raise HTTPException(status_code=404, detail=f"Rasterized page missing for page {page.page_number}")

        tokens = extract_tokens_with_bboxes(image_path)
        lines_by_key: dict[tuple[int, int, int], list[dict[str, Any]]] = {}
        for token_index, token in enumerate(tokens):
            key = (
                int(token.get("block_num", 0)),
                int(token.get("paragraph_num", 0)),
                int(token.get("line_num", 0)),
            )
            lines_by_key.setdefault(key, []).append({**token, "token_index": token_index})

        lines: list[dict[str, Any]] = []
        for line_index, (line_key, line_tokens) in enumerate(lines_by_key.items()):
            lines.append(
                {
                    "id": f"{page.page_number}-{line_index}",
                    "block_num": line_key[0],
                    "paragraph_num": line_key[1],
                    "line_num": line_key[2],
                    "text": " ".join(token["token"] for token in line_tokens),
                    "tokens": line_tokens,
                }
            )

        page_payloads.append(
            {
                "page_id": page.id,
                "page_number": page.page_number,
                "image_url": f"/api/v1/documents/{document.id}/pages/{page.page_number}/image",
                "text": "\n".join(line["text"] for line in lines),
                "lines": lines,
                "tokens": tokens,
            }
        )

    return {"document_id": document.id, "document_type": document.document_type, "pages": page_payloads}


@router.patch("/documents/{document_id}/status")
def update_document_status(document_id: int, status: str, db: Session = Depends(get_db), user_id: CurrentUserId = None) -> dict[str, Any]:
    document = db.query(DocumentRecord).filter(DocumentRecord.id == document_id, owner_scope(DocumentRecord.owner_id, user_id)).first()
    if document is None:
        raise HTTPException(status_code=404, detail="Document not found")

    try:
        _transition_document_status(document, status)
    except ValueError as exc:
        raise HTTPException(status_code=400, detail=str(exc)) from exc
    db.commit()
    db.refresh(document)
    return {
        "id": document.id,
        "status": document.status,
    }


@router.get("/documents/{document_id}/pages/{page_number}/image")
def get_document_page_image(
    document_id: int,
    page_number: int,
    db: Session = Depends(get_db),
    user_id: CurrentUserId = None,
) -> Response:
    document = db.query(DocumentRecord).filter(DocumentRecord.id == document_id, owner_scope(DocumentRecord.owner_id, user_id)).first()
    if document is None:
        raise HTTPException(status_code=404, detail="Document not found")
    page = next((item for item in document.pages if item.page_number == page_number), None)
    if page is None:
        raise HTTPException(status_code=404, detail="Page not found")
    if storage_enabled() and page.storage_key:
        return Response(content=download_file(page.storage_key), media_type="image/png")
    image_path = Path(page.image_path)
    if not image_path.exists():
        image_path = RAW_DATASET_ROOT / str(document_id) / f"page-{page_number:03d}.png"
    if not image_path.exists():
        raise HTTPException(status_code=404, detail="Page image not found")
    return Response(content=image_path.read_bytes(), media_type="image/png")


def _cleanup_artifacts(page_paths: list[Path], storage_keys: list[str]) -> None:
    for page_path in page_paths:
        try:
            page_path.unlink(missing_ok=True)
        except OSError:
            logger.warning("Unable to remove local processing artifact %s", page_path, exc_info=True)
    if storage_enabled():
        for storage_key in storage_keys:
            try:
                delete_file(storage_key)
            except Exception:
                logger.warning("Unable to remove storage processing artifact %s", storage_key, exc_info=True)


@celery_app.task(name="app.ingestion_pipeline.process_document_job")
def process_document_job(document_id: int, source_key: str, original_name: str, owner_id: str | None = None) -> dict[str, Any]:
    db = SessionLocal()
    set_rls_user(db, owner_id)
    source_temp_dir: tempfile.TemporaryDirectory[str] | None = None

    def _materialize_source(key: str) -> Path:
        nonlocal source_temp_dir
        if storage_enabled():
            source_temp_dir = tempfile.TemporaryDirectory(prefix=f"ocr-source-{document_id}-")
            path = Path(source_temp_dir.name) / Path(original_name).name
            path.write_bytes(download_file(key))
            return path

        legacy_path = Path(key.replace("\\", "/"))
        if legacy_path.is_absolute():
            return legacy_path
        return UPLOAD_ROOT / str(document_id) / f"source{Path(original_name).suffix.lower()}"

    page_paths: list[Path] = []
    new_storage_keys: list[str] = []
    new_page_ids: list[int] = []
    processing_artifact_id = f"{document_id}/runs/{uuid.uuid4().hex}"
    final_document_type = "unknown"
    processing_started_at: datetime | None = None
    processing_start_time: float | None = None
    previous_status: str | None = None
    previous_timing: tuple[Any, Any, Any, int, str] = (None, None, None, 0, "unknown")

    try:
        document = db.query(DocumentRecord).filter(DocumentRecord.id == document_id, DocumentRecord.owner_id == owner_id).first()
        if document is None:
            raise FileNotFoundError(f"Document record missing: {document_id}")
        if document.status == "CANCELED":
            return {
                "document_id": document_id,
                "original_name": original_name,
                "document_type": str(document.document_type or "unknown"),
                "page_count": document.page_count,
                "pages": [],
                "status": "canceled",
            }

        previous_status = document.status
        previous_page_rows = list(document.pages)
        previous_attribute_ids = [attribute.id for attribute in document.attributes]
        previous_page_ids = [page.id for page in previous_page_rows]
        previous_page_paths = [Path(page.image_path) for page in previous_page_rows]
        previous_storage_keys = [page.storage_key for page in previous_page_rows if page.storage_key]
        previous_timing = (
            document.processing_started_at,
            document.processing_finished_at,
            document.processing_duration_seconds,
            document.page_count,
            document.document_type,
        )

        processing_started_at = datetime.now(timezone.utc)
        processing_start_time = time.perf_counter()
        document.processing_started_at = processing_started_at
        document.processing_finished_at = None
        document.processing_duration_seconds = None
        _transition_document_status(document, "PROCESSING")
        db.commit()

        path = _materialize_source(source_key)
        if not path.exists():
            logger.error("Source file missing for %s: %s", document_id, source_key)
            raise FileNotFoundError(f"Source file missing: {source_key}")

        extension = path.suffix.lower()
        if extension == ".pdf":
            page_paths = _rasterize_pdf(path, processing_artifact_id)
        elif extension in {".docx", ".doc"}:
            page_paths = _rasterize_office_document(path, processing_artifact_id)
        elif extension in IMAGE_EXTENSIONS:
            page_paths = _normalize_image_document(path, processing_artifact_id)
        else:
            raise ValueError(f"Unsupported format for processing: {extension}")

        detected_document_type = _detect_document_type(page_paths, original_name)
        document.document_type = detected_document_type["document_type"]
        logger.info(
            "Detected document type for %s: %s (confidence=%.3f)",
            original_name,
            detected_document_type["document_type"],
            float(detected_document_type.get("confidence", 0.0)),
        )

        staged_attributes: list[dict[str, Any]] = []

        for index, page_path in enumerate(page_paths, start=1):
            page = Page(
                document_id=document_id,
                page_number=index,
                image_path=str(page_path),
                storage_key=f"{document.owner_id}/{processing_artifact_id}/pages/page-{index:03d}.png",
            )
            db.add(page)
            db.flush()
            new_page_ids.append(page.id)
            new_storage_keys.append(page.storage_key)
            upload_file(page_path, page.storage_key)

            attribute_payloads = _build_attribute_payloads_for_page(
                page.id,
                document_id,
                page_path,
                document_type=document.document_type,
            )
            for payload in attribute_payloads:
                staged_attributes.append(
                    {
                        "temp_id": str(uuid.uuid4()),
                        "page_id": page.id,
                        "page_number": index,
                        "document_id": document_id,
                        "entity_type_label": payload["entity_type_label"],
                        "extracted_value": payload["extracted_value"],
                        "original_prediction": payload["extracted_value"],
                        "bounding_boxes": payload["bounding_boxes"],
                        "confidence_score": payload["confidence_score"],
                        "validation_status": payload["validation_status"],
                        "source": "model",
                        "saved": False,
                    }
                )

        db.expire(document)
        db.refresh(document)
        if document.status == "CANCELED":
            db.rollback()
            _cleanup_artifacts(page_paths, new_storage_keys)
            return {
                "document_id": document_id,
                "original_name": original_name,
                "document_type": str(document.document_type or "unknown"),
                "page_count": document.page_count,
                "pages": [],
                "status": "canceled",
            }

        document.page_count = len(page_paths)
        document.extracted_attributes = staged_attributes
        document.processing_finished_at = datetime.now(timezone.utc)
        document.processing_duration_seconds = time.perf_counter() - processing_start_time
        _transition_document_status(document, "PROCESSED")
        final_document_type = str(document.document_type or "unknown")
        if previous_attribute_ids:
            db.query(Attribute).filter(Attribute.id.in_(previous_attribute_ids)).delete(synchronize_session=False)
        if previous_page_ids:
            db.query(Page).filter(Page.id.in_(previous_page_ids)).delete(synchronize_session=False)
        db.commit()
        _cleanup_artifacts(previous_page_paths, previous_storage_keys)
    except Exception:
        db.rollback()
        _cleanup_artifacts(page_paths, new_storage_keys)
        document = db.get(DocumentRecord, document_id)
        if document is not None:
            if document.status == "CANCELED":
                pass
            elif previous_status in {"PROCESSED", "APPROVED"}:
                (
                    document.processing_started_at,
                    document.processing_finished_at,
                    document.processing_duration_seconds,
                    document.page_count,
                    document.document_type,
                ) = previous_timing
                document.status = previous_status
            elif processing_started_at is not None and processing_start_time is not None:
                document.processing_finished_at = datetime.now(timezone.utc)
                document.processing_duration_seconds = time.perf_counter() - processing_start_time
                _transition_document_status(document, "FAILED")
            db.commit()
        logger.exception("Document processing failed for %s (%s)", document_id, original_name)
        raise
    finally:
        if source_temp_dir is not None:
            source_temp_dir.cleanup()
        db.close()

    return {
        "document_id": document_id,
        "original_name": original_name,
        "document_type": final_document_type,
        "page_count": len(page_paths),
        "pages": [str(page) for page in page_paths],
        "status": "processed",
    }


def _store_raw_dataset_pages(page_paths: list[Path], document_id: str | int) -> list[Path]:
    dataset_dir = RAW_DATASET_ROOT / str(document_id)
    dataset_dir.mkdir(parents=True, exist_ok=True)
    stored_paths: list[Path] = []
    for page_number, page_path in enumerate(page_paths, start=1):
        output_path = dataset_dir / f"page-{page_number:03d}.png"
        shutil.copy2(page_path, output_path)
        stored_paths.append(output_path)
    return stored_paths


def _detect_document_type(page_paths: list[Path], original_name: str) -> dict[str, Any]:
    if not page_paths:
        return {"document_type": "unknown", "confidence": 0.0, "scores": {}}

    first_page_tokens = extract_tokens_with_bboxes(page_paths[0])
    text = " ".join(token["token"] for token in first_page_tokens)
    line_count = len(
        {
            (
                int(token.get("block_num", 0)),
                int(token.get("paragraph_num", 0)),
                int(token.get("line_num", 0)),
            )
            for token in first_page_tokens
        }
    )
    return classify_document_type(
        text=text,
        filename=original_name,
        layout_hints={
            "page_count": len(page_paths),
            "token_count": len(first_page_tokens),
            "line_count": line_count,
        },
    )


def _rasterize_pdf(path: Path, document_id: str) -> List[Path]:
    asset_dir = RAW_DATASET_ROOT / str(document_id)
    asset_dir.mkdir(parents=True, exist_ok=True)
    try:
        pdf_document = fitz.open(str(path))
    except Exception as exc:
        logger.exception("PDF rasterization failed for %s", path)
        raise RuntimeError(f"Failed to rasterize PDF document: {path}") from exc

    page_paths: List[Path] = []
    try:
        for index, page in enumerate(pdf_document, start=1):
            output_path = asset_dir / f"page-{index:03d}.png"
            matrix = fitz.Matrix(300 / 72, 300 / 72)
            pixmap = page.get_pixmap(matrix=matrix, alpha=False)
            pixmap.save(str(output_path))
            page_paths.append(output_path)
    finally:
        pdf_document.close()
    return page_paths


def _rasterize_office_document(path: Path, document_id: str) -> List[Path]:
    asset_dir = RAW_DATASET_ROOT / str(document_id)
    asset_dir.mkdir(parents=True, exist_ok=True)
    with tempfile.TemporaryDirectory() as temp_dir:
        temp_dir_path = Path(temp_dir)
        pdf_path = temp_dir_path / f"{path.stem}.pdf"
        try:
            _convert_with_libreoffice(path, temp_dir_path)
            if not pdf_path.exists():
                generated_pdfs = list(temp_dir_path.glob("*.pdf"))
                if not generated_pdfs:
                    raise RuntimeError("LibreOffice did not emit a PDF file")
                pdf_path = generated_pdfs[0]
            return _rasterize_pdf(pdf_path, document_id)
        except Exception as exc:
            if path.suffix.lower() == ".docx":
                logger.warning("Falling back to python-docx rendering for %s: %s", path, exc)
                try:
                    return _rasterize_docx_fallback(path, document_id)
                except Exception as fallback_exc:
                    logger.exception("DOCX fallback rasterization failed for %s", path)
                    raise RuntimeError(f"Failed to rasterize office document: {path}") from fallback_exc
            logger.exception("Office document rasterization failed for %s", path)
            raise RuntimeError(f"Failed to rasterize office document: {path}") from exc


def _convert_with_libreoffice(source_path: Path, output_dir: Path) -> None:
    command = [
        SOFFICE_CMD,
        "--headless",
        "--convert-to",
        "pdf",
        "--outdir",
        str(output_dir),
        str(source_path),
    ]
    try:
        completed = subprocess.run(command, check=True, capture_output=True, text=True)
        if completed.stdout:
            logger.info("LibreOffice output for %s: %s", source_path.name, completed.stdout.strip())
        if completed.stderr:
            logger.warning("LibreOffice warnings for %s: %s", source_path.name, completed.stderr.strip())
    except FileNotFoundError as exc:
        raise RuntimeError(
            "LibreOffice is required to rasterize DOC files but was not found on the PATH"
        ) from exc


def _rasterize_docx_fallback(path: Path, document_id: str) -> List[Path]:
    document = DocxDocument(path)
    asset_dir = RAW_DATASET_ROOT / str(document_id)
    asset_dir.mkdir(parents=True, exist_ok=True)
    renderer = _DocxFallbackRenderer(document, asset_dir)
    return renderer.render()


class _DocxFallbackRenderer:
    page_size = (1600, 2200)
    margin = 100
    line_height = 34

    def __init__(self, document: DocxDocumentPart, asset_dir: Path) -> None:
        self.document = document
        self.asset_dir = asset_dir
        self.pages: list[Image.Image] = []
        self._new_page()

    def render(self) -> List[Path]:
        for block in self._iter_blocks(self.document):
            if isinstance(block, Paragraph):
                self._render_paragraph(block)
            else:
                self._render_table(block)
        self._draw_header_footer(self.pages[-1])
        paths = []
        for index, page in enumerate(self.pages, start=1):
            output_path = self.asset_dir / f"page-{index:03d}.png"
            page.save(output_path, format="PNG")
            paths.append(output_path)
        return paths

    def _new_page(self) -> None:
        if self.pages:
            self._draw_header_footer(self.pages[-1])
        self.pages.append(Image.new("RGB", self.page_size, "white"))
        self.cursor_y = self.margin + 80

    def _available_height(self) -> int:
        return self.page_size[1] - self.margin - 80

    def _ensure_space(self, height: int) -> None:
        if self.cursor_y + height > self._available_height():
            self._new_page()

    def _render_paragraph(self, paragraph: Paragraph) -> None:
        text = paragraph.text
        font_size = self._paragraph_font_size(paragraph)
        font = self._font(font_size)
        line_height = max(self.line_height, font_size + 12)
        available_width = self.page_size[0] - (2 * self.margin)
        lines = self._wrap_text(text, font, available_width)
        if not lines:
            lines = [""]
        self._ensure_space(len(lines) * line_height)
        draw = ImageDraw.Draw(self.pages[-1])
        x = self.margin
        for line in lines:
            draw.text((x, self.cursor_y), line, fill="black", font=font)
            self.cursor_y += line_height
        self._render_paragraph_images(paragraph)
        if self._has_page_break(paragraph):
            self._new_page()

    def _render_table(self, table: Table) -> None:
        rows = []
        for row in table.rows:
            rows.append([" ".join(cell.text.split()) for cell in row.cells])
        if not rows:
            return
        columns = max(len(row) for row in rows)
        column_width = (self.page_size[0] - 2 * self.margin) // max(columns, 1)
        font = self._font(24)
        row_height = 48
        for row in rows:
            self._ensure_space(row_height)
            draw = ImageDraw.Draw(self.pages[-1])
            for column in range(columns):
                left = self.margin + column * column_width
                top = self.cursor_y
                draw.rectangle((left, top, left + column_width, top + row_height), outline="black")
                value = row[column] if column < len(row) else ""
                draw.text((left + 8, top + 10), self._truncate_text(value, font, column_width - 16), fill="black", font=font)
            self.cursor_y += row_height

    def _render_paragraph_images(self, paragraph: Paragraph) -> None:
        for blip in paragraph._p.xpath(".//a:blip"):
            relation_id = blip.get(qn("r:embed"))
            if not relation_id:
                continue
            image_part = self.document.part.related_parts[relation_id]
            with Image.open(BytesIO(image_part.blob)) as source:
                image = source.convert("RGB")
            max_width = self.page_size[0] - 2 * self.margin
            if image.width > max_width:
                ratio = max_width / image.width
                image = image.resize((max_width, max(1, int(image.height * ratio))))
            self._ensure_space(image.height + 12)
            self.pages[-1].paste(image, (self.margin, self.cursor_y))
            self.cursor_y += image.height + 12

    def _draw_header_footer(self, page: Image.Image) -> None:
        draw = ImageDraw.Draw(page)
        font = self._font(22)
        header = " ".join(paragraph.text.strip() for paragraph in self.document.sections[0].header.paragraphs if paragraph.text.strip())
        footer = " ".join(paragraph.text.strip() for paragraph in self.document.sections[0].footer.paragraphs if paragraph.text.strip())
        if header:
            draw.text((self.margin, self.margin), header, fill="black", font=font)
        if footer:
            draw.text((self.margin, self.page_size[1] - self.margin - 30), footer, fill="black", font=font)

    @staticmethod
    def _iter_blocks(parent: DocxDocumentPart):
        body = parent.element.body
        for child in body.iterchildren():
            if child.tag == qn("w:p"):
                yield Paragraph(child, parent)
            elif child.tag == qn("w:tbl"):
                yield Table(child, parent)

    @staticmethod
    def _paragraph_font_size(paragraph: Paragraph) -> int:
        for run in paragraph.runs:
            if run.font.size:
                return max(12, int(run.font.size.pt))
        style_size = paragraph.style.font.size if paragraph.style and paragraph.style.font else None
        return max(12, int(style_size.pt)) if style_size else 24

    @staticmethod
    def _font(size: int) -> ImageFont.ImageFont:
        try:
            return ImageFont.truetype("DejaVuSans.ttf", size)
        except OSError:
            return ImageFont.load_default()

    @staticmethod
    def _wrap_text(text: str, font: ImageFont.ImageFont, width: int) -> list[str]:
        words = text.split()
        if not words:
            return []
        lines = []
        current = words[0]
        for word in words[1:]:
            candidate = f"{current} {word}"
            if font.getlength(candidate) <= width:
                current = candidate
            else:
                lines.append(current)
                current = word
        lines.append(current)
        return lines

    @staticmethod
    def _truncate_text(text: str, font: ImageFont.ImageFont, width: int) -> str:
        while text and font.getlength(text) > width:
            text = text[:-1]
        return text

    @staticmethod
    def _has_page_break(paragraph: Paragraph) -> bool:
        return bool(paragraph._p.xpath('.//w:br[@w:type="page"]')) or bool(
            paragraph._p.xpath('.//w:lastRenderedPageBreak')
        )


def _normalize_image_document(path: Path, document_id: str) -> List[Path]:
    asset_dir = RAW_DATASET_ROOT / str(document_id)
    asset_dir.mkdir(parents=True, exist_ok=True)
    page_paths: List[Path] = []

    with Image.open(path) as image:
        frames = list(ImageSequence.Iterator(image)) if getattr(image, "is_animated", False) else [image]
        for index, frame in enumerate(frames, start=1):
            output_path = asset_dir / f"page-{index:03d}.png"
            frame.convert("RGB").save(output_path, format="PNG")
            page_paths.append(output_path)

    return page_paths


def _build_attribute_payloads_for_page(
    page_id: int,
    document_id: int,
    image_path: Path,
    document_type: str | None = None,
) -> List[dict[str, Any]]:
    tokens = extract_tokens_with_bboxes(image_path)
    if not tokens:
        return []

    token_predictions: List[dict[str, Any]] = []
    try:
        inference_service = LayoutLMv3InferenceService()
        inference_result = inference_service.infer_page(
            image_path=image_path,
            tokens=[item["token"] for item in tokens],
            normalized_bboxes=[item["bbox"] for item in tokens],
            document_type=document_type,
        )
        token_predictions = inference_result.get("tokens", [])
    except Exception as exc:
        logger.warning("LayoutLMv3 inference failed for %s; using OCR fallback: %s", image_path, exc)
        token_predictions = [
            {
                "token": item["token"],
                "label": "OCR",
                "confidence": max(0.0, min(1.0, float(item["confidence"]) / 100.0)),
                "bbox": item["bbox"],
            }
            for item in tokens
        ]

    return build_attribute_payloads(
        page_id=page_id,
        document_id=document_id,
        token_predictions=token_predictions,
        document_type=document_type,
        threshold=0.75,
    )
